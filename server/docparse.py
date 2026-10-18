"""文档解析：PDF / DOCX / XLSX / PPTX / EPUB / 图片OCR → 纯文本。

设计约定（与 tools.py 一致）：
- 全部函数返回 str；失败返回 "[error] 可读原因"（喂给小模型也能看懂）。
- 解析库按需 import（懒加载），避免拖慢启动、避免无谓内存占用。
- 输出带 "--- 第N页 ---" 结构标记，便于 RAG 分块与引用定位。

被三处共用：
- server/rag.py   ：全库索引（文档/图片也入库）
- server/mcp_server.py ：local_doc_ask 单文档问答
- chatui/chat_app.py   ：拖拽文档附件（打包时经 spec pathex 收录）
"""
from __future__ import annotations

import os
import re
import zipfile
from pathlib import Path

SUPPORTED_EXTS = {".pdf", ".docx", ".xlsx", ".pptx", ".epub",
                  ".txt", ".md", ".log", ".py", ".json", ".yaml", ".yml", ".csv"}
DOC_EXTS = {".pdf", ".docx", ".xlsx", ".pptx", ".epub"}       # 需要解析的"文档"
OCR_EXTS = {".png", ".jpg", ".jpeg", ".bmp", ".webp"}

MAX_DOC_BYTES = 50 * 1024 * 1024     # 单文档上限
MAX_TEXT_CHARS = 200_000             # 单文档输出上限（防 OOM / prompt 爆炸）
MAX_XLSX_ROWS = 500                  # 每 sheet 行数上限
MAX_PDF_OCR_PAGES = 10               # 扫描 PDF 最多 OCR 前N页
MIN_PDF_TEXT_CHARS = 50              # 低于此值视为扫描件，走 OCR
MAX_OCR_IMAGE_BYTES = 5 * 1024 * 1024

_TAG_RE = re.compile(r"<[^>]+>")


def supported(path) -> bool:
    return Path(path).suffix.lower() in (SUPPORTED_EXTS | OCR_EXTS)


def extract_text(path, ocr_fallback: bool = True) -> str:
    """按扩展名解析为纯文本；OCR 图片也走这里。"""
    p = Path(path)
    ext = p.suffix.lower()
    if ext not in (SUPPORTED_EXTS | OCR_EXTS):
        return f"[error] 不支持的文件类型: {ext or '(无后缀)'}"
    if not p.is_file():
        return f"[error] 文件不存在: {path}"
    if p.stat().st_size > MAX_DOC_BYTES:
        return f"[error] 文件过大（>{MAX_DOC_BYTES >> 20}MB）: {p.name}"

    try:
        if ext == ".pdf":
            text = _pdf(p, ocr_fallback)
        elif ext == ".docx":
            text = _docx(p)
        elif ext == ".xlsx":
            text = _xlsx(p)
        elif ext == ".pptx":
            text = _pptx(p)
        elif ext == ".epub":
            text = _epub(p)
        elif ext in OCR_EXTS:
            text = ocr_image(p)
        else:                                    # 纯文本类
            text = p.read_text(encoding="utf-8", errors="replace")
    except ImportError as e:
        return f"[error] 缺少解析库: {e}（pip 安装对应依赖后重试）"
    except Exception as e:
        return f"[error] 解析失败 {p.name}: {e}"

    text = (text or "").strip()
    if not text:
        return f"[error] 未提取到文本内容: {p.name}（扫描件或空文档）"
    if len(text) > MAX_TEXT_CHARS:
        text = text[:MAX_TEXT_CHARS] + f"\n…（已截断，原文 {len(text)} 字符）"
    return text


# ---------- PDF ----------
def _pdf(p: Path, ocr_fallback: bool) -> str:
    from pypdf import PdfReader
    pages = []
    # 用 with 显式关文件：PdfReader 不会自己关句柄，Windows 下会锁住用户文件
    with open(p, "rb") as f:
        reader = PdfReader(f)
        for i, page in enumerate(reader.pages, 1):
            try:
                t = (page.extract_text() or "").strip()
            except Exception:
                t = ""
            pages.append(f"--- 第{i}页 ---\n{t}" if t else f"--- 第{i}页 ---")
    text = "\n\n".join(pages)
    total = len(re.sub(r"[\s\-]|第\d+页", "", text))
    if total < MIN_PDF_TEXT_CHARS and ocr_fallback:
        ocr_text = _pdf_ocr(p)
        if not ocr_text.startswith("[error]"):
            text = f"{text}\n\n[OCR 提取]\n{ocr_text}"
    return text


def _pdf_ocr(p: Path) -> str:
    """扫描 PDF：pypdfium2 渲页 → RapidOCR（不经 Pillow，直接 to_numpy）。"""
    try:
        import pypdfium2 as pdfium
    except ImportError:
        return "[error] 扫描PDF需要 pypdfium2"
    ocr = _get_ocr()
    if ocr is None:
        return "[error] OCR 不可用"
    pdf = pdfium.PdfDocument(str(p))
    out = []
    try:
        for i in range(min(len(pdf), MAX_PDF_OCR_PAGES)):
            bitmap = None
            try:
                bitmap = pdf[i].render(scale=2.0)
                img = bitmap.to_numpy()          # HWC, RGB(实际BGR皆可)
                result, _ = ocr(img)
                lines = [r[1] for r in (result or [])]
                if lines:
                    out.append(f"--- 第{i + 1}页(OCR) ---\n" + "\n".join(lines))
            except Exception:
                continue
            finally:
                # 显式释放位图：不关会一直占着 PDF 句柄，Windows 下用户就删不掉这个文件
                if bitmap is not None:
                    try:
                        bitmap.close()
                    except Exception:
                        pass
    finally:
        try:
            pdf.close()
        except Exception:
            pass
    return "\n\n".join(out) if out else "[error] OCR 未提取到文字"


# ---------- DOCX ----------
def _docx(p: Path) -> str:
    from docx import Document
    doc = Document(str(p))
    parts = []
    for para in doc.paragraphs:
        t = para.text.strip()
        if t:
            parts.append(t)
    for ti, table in enumerate(doc.tables, 1):
        rows = []
        for row in table.rows:
            cells = [c.text.strip().replace("\n", " ") for c in row.cells]
            rows.append(" | ".join(cells))
        if rows:
            parts.append(f"[表格{ti}]\n" + "\n".join(rows))
    return "\n".join(parts)


# ---------- XLSX ----------
def _xlsx(p: Path) -> str:
    import openpyxl
    wb = openpyxl.load_workbook(str(p), read_only=True, data_only=True)
    parts = []
    for ws in wb.worksheets:
        rows = []
        for ri, row in enumerate(ws.iter_rows(values_only=True), 1):
            if ri > MAX_XLSX_ROWS:
                rows.append(f"…（工作表 {ws.title} 仅保留前 {MAX_XLSX_ROWS} 行）")
                break
            cells = ["" if v is None else str(v).replace("\n", " ") for v in row]
            if any(c.strip() for c in cells):
                rows.append(" | ".join(cells))
        if rows:
            parts.append(f"=== 工作表: {ws.title} ===\n" + "\n".join(rows))
    wb.close()
    return "\n".join(parts)


# ---------- PPTX ----------
def _pptx(p: Path) -> str:
    from pptx import Presentation
    prs = Presentation(str(p))
    parts = []
    for i, slide in enumerate(prs.slides, 1):
        lines = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                t = shape.text_frame.text.strip()
                if t:
                    lines.append(t)
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    cells = [c.text.strip() for c in row.cells]
                    lines.append(" | ".join(cells))
        if lines:
            parts.append(f"--- 幻灯片{i} ---\n" + "\n".join(lines))
    return "\n".join(parts)


# ---------- EPUB ----------
def _epub(p: Path) -> str:
    parts = []
    with zipfile.ZipFile(str(p)) as z:
        for name in z.namelist():
            if name.lower().endswith((".xhtml", ".html", ".htm", ".txt")) and "css" not in name:
                try:
                    raw = z.read(name).decode("utf-8", errors="replace")
                except Exception:
                    continue
                raw = re.sub(r"(?is)<(script|style).*?</\1>", "", raw)
                txt = _TAG_RE.sub(" ", raw)
                txt = re.sub(r"[ \t]+", " ", txt)
                txt = "\n".join(x.strip() for x in txt.splitlines() if x.strip())
                if len(txt) > 80:
                    parts.append(txt)
    return "\n".join(parts[:80])       # epub 章节数上限


# ---------- OCR ----------
_OCR = None


def _get_ocr():
    global _OCR
    if _OCR is None:
        try:
            from rapidocr_onnxruntime import RapidOCR
            _OCR = RapidOCR()
        except Exception:
            _OCR = False
    return _OCR or None


def ocr_image(path) -> str:
    """图片文字提取；无文字返回空串，库缺失返回 [error]。"""
    p = Path(path)
    if p.suffix.lower() not in OCR_EXTS:
        return "[error] 非 OCR 支持的图片类型"
    if p.stat().st_size > MAX_OCR_IMAGE_BYTES:
        return f"[error] 图片过大（>{MAX_OCR_IMAGE_BYTES >> 20}MB）"
    ocr = _get_ocr()
    if ocr is None:
        return "[error] OCR 组件不可用（缺 rapidocr_onnxruntime）"
    try:
        result, _ = ocr(str(p))
        lines = [r[1] for r in (result or [])]
        return "\n".join(lines)
    except Exception as e:
        return f"[error] OCR 失败: {e}"
